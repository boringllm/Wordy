"""Generate examples/sample_template.docx — a corporate-style template with a
cover page, custom heading colors, header/footer with placeholders, and a
styled table. Used to test/demo the MCP server."""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "sample_template.docx")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # dark corporate blue
ACCENT2 = RGBColor(0xC0, 0x50, 0x4D)  # red accent

doc = Document()

# ----- restyle built-in styles -----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

for lvl, size, color in ((1, 20, ACCENT), (2, 15, ACCENT), (3, 12, ACCENT2)):
    st = doc.styles["Heading %d" % lvl]
    st.font.name = "Georgia"
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True

title = doc.styles["Title"]
title.font.name = "Georgia"
title.font.size = Pt(36)
title.font.color.rgb = ACCENT

quote = doc.styles["Quote"]
quote.font.italic = True
quote.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

# ----- header / footer with placeholders -----
section = doc.sections[0]
section.header.paragraphs[0].text = "{{company}} — Confidential"
section.footer.paragraphs[0].text = "{{title}} | Page"
section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# ----- cover page -----
p = doc.add_paragraph(style="Title")
p.add_run("{{title}}")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("{{subtitle}}")
r.font.size = Pt(16)
r.font.color.rgb = ACCENT2

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Prepared by {{author}} — {{date}}").italic = True

# page break ends the cover page
p = doc.add_paragraph()
p.add_run().add_break(WD_BREAK.PAGE)

# ----- sample body content (replaced on generation) -----
doc.add_paragraph("Sample heading", style="Heading 1")
doc.add_paragraph("This body content is placeholder material that generated "
                  "reports replace.", style="Normal")
t = doc.add_table(rows=2, cols=2, style="Light Grid Accent 1")
t.cell(0, 0).text = "Sample"
t.cell(0, 1).text = "Table"
t.cell(1, 0).text = "A"
t.cell(1, 1).text = "B"

doc.save(OUT)
print("wrote", OUT)
