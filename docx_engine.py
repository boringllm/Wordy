"""
docx_engine.py — Core engine for the Word Template MCP server.

Three capabilities, template-agnostic:
  1. inspect_template(path)  -> full JSON description of a .docx template
     (styles, fonts, colors, theme, sections, headers/footers, body blocks,
      placeholders, and a suggested markdown->style mapping).
  2. create_document(...)    -> render a markdown report into a new .docx that
     strictly reuses the template's styles, theme, headers/footers, cover page.
  3. validate_document(...)  -> verify a produced .docx conforms to a template.

The design principle: the output document is the template file itself
(opened, body selectively kept/cleared, content appended), so every style,
numbering definition, theme, header, footer, page setup and cover element
is preserved byte-for-byte. Markdown content is then mapped onto the
template's own named styles.
"""

from __future__ import annotations

import os
import re
import copy
import zipfile
from typing import Any, Optional

import mistune
from docx import Document
from docx.shared import Pt, RGBColor, Emu
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")

DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------


def _safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default


def _color_hex(color_obj) -> Optional[str]:
    try:
        if color_obj is None or color_obj.rgb is None:
            return None
        return "#%s" % str(color_obj.rgb)
    except Exception:
        return None


def _font_info(font) -> dict:
    """Serialize a python-docx Font object (values are None when inherited)."""
    return {
        "name": _safe(lambda: font.name),
        "size_pt": _safe(lambda: font.size.pt if font.size else None),
        "bold": _safe(lambda: font.bold),
        "italic": _safe(lambda: font.italic),
        "underline": _safe(lambda: font.underline),
        "all_caps": _safe(lambda: font.all_caps),
        "small_caps": _safe(lambda: font.small_caps),
        "color": _safe(lambda: _color_hex(font.color)),
    }


def _iter_body_blocks(doc) -> list:
    """Return top-level body block elements (w:p and w:tbl), in order."""
    body = doc.element.body
    return [el for el in body.iterchildren() if el.tag in (qn("w:p"), qn("w:tbl"))]


def _paragraph_text(p_el) -> str:
    return "".join(node.text or "" for node in p_el.iter(qn("w:t")))


def _has_page_break(p_el) -> bool:
    for br in p_el.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _has_section_break(p_el) -> bool:
    return p_el.find(qn("w:pPr")) is not None and \
        p_el.find(qn("w:pPr")).find(qn("w:sectPr")) is not None


def _para_style_name(p_el, doc) -> str:
    pPr = p_el.find(qn("w:pPr"))
    if pPr is not None:
        pStyle = pPr.find(qn("w:pStyle"))
        if pStyle is not None:
            sid = pStyle.get(qn("w:val"))
            for s in doc.styles:
                if s.style_id == sid:
                    return s.name
            return sid
    return "Normal"


def _style_index(doc) -> dict:
    """Case-insensitive lookup: style name AND style_id -> style object."""
    idx = {}
    for s in doc.styles:
        try:
            idx[s.name.lower()] = s
            idx[s.style_id.lower()] = s
        except Exception:
            continue
    return idx


# ---------------------------------------------------------------------------
# 1. TEMPLATE INSPECTION
# ---------------------------------------------------------------------------


def _theme_info(path: str) -> dict:
    """Parse word/theme/theme1.xml for theme colors and major/minor fonts."""
    out: dict = {"colors": {}, "fonts": {}}
    try:
        from lxml import etree
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if n.startswith("word/theme/")]
            if not names:
                return out
            root = etree.fromstring(z.read(names[0]))
        ns = {"a": DRAWINGML_NS}
        scheme = root.find(".//a:clrScheme", ns)
        if scheme is not None:
            for child in scheme:
                tag = etree.QName(child).localname
                srgb = child.find("a:srgbClr", ns)
                sysc = child.find("a:sysClr", ns)
                if srgb is not None:
                    out["colors"][tag] = "#" + srgb.get("val")
                elif sysc is not None and sysc.get("lastClr"):
                    out["colors"][tag] = "#" + sysc.get("lastClr")
        for kind in ("majorFont", "minorFont"):
            el = root.find(".//a:%s/a:latin" % kind, ns)
            if el is not None:
                out["fonts"][kind] = el.get("typeface")
    except Exception as e:
        out["error"] = str(e)
    return out


def _doc_defaults(doc) -> dict:
    out = {}
    try:
        normal = doc.styles["Normal"]
        out["normal_style_font"] = _font_info(normal.font)
    except Exception:
        pass
    return out


def _collect_placeholders(doc) -> list:
    found = set()

    def scan_paragraphs(paragraphs):
        for p in paragraphs:
            for m in PLACEHOLDER_RE.finditer(p.text):
                found.add(m.group(1))

    def scan_tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    scan_paragraphs(cell.paragraphs)
                    scan_tables(cell.tables)

    scan_paragraphs(doc.paragraphs)
    scan_tables(doc.tables)
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part is not None:
                scan_paragraphs(part.paragraphs)
                scan_tables(part.tables)
    return sorted(found)


def _sections_info(doc) -> list:
    sections = []
    for i, s in enumerate(doc.sections):
        def hf_text(part):
            if part is None:
                return None
            txt = "\n".join(p.text for p in part.paragraphs if p.text.strip())
            return txt or None
        sections.append({
            "index": i,
            "page_width_cm": _safe(lambda: round(s.page_width.cm, 2)),
            "page_height_cm": _safe(lambda: round(s.page_height.cm, 2)),
            "orientation": _safe(lambda: str(s.orientation).split(".")[-1].split(" ")[0]),
            "margins_cm": {
                "top": _safe(lambda: round(s.top_margin.cm, 2)),
                "bottom": _safe(lambda: round(s.bottom_margin.cm, 2)),
                "left": _safe(lambda: round(s.left_margin.cm, 2)),
                "right": _safe(lambda: round(s.right_margin.cm, 2)),
            },
            "different_first_page": _safe(lambda: s.different_first_page_header_footer),
            "header_text": hf_text(_safe(lambda: s.header)),
            "footer_text": hf_text(_safe(lambda: s.footer)),
            "first_page_header_text": hf_text(_safe(lambda: s.first_page_header)),
            "first_page_footer_text": hf_text(_safe(lambda: s.first_page_footer)),
        })
    return sections


def _body_outline(doc, preview_chars: int = 100) -> list:
    blocks = []
    for i, el in enumerate(_iter_body_blocks(doc)):
        if el.tag == qn("w:tbl"):
            rows = len(el.findall(qn("w:tr")))
            style_el = el.find(qn("w:tblPr") + "/" + qn("w:tblStyle"))
            blocks.append({
                "index": i, "type": "table", "rows": rows,
                "style": style_el.get(qn("w:val")) if style_el is not None else None,
                "preview": _paragraph_text(el)[:preview_chars],
            })
        else:
            text = _paragraph_text(el)
            info = {
                "index": i, "type": "paragraph",
                "style": _para_style_name(el, doc),
                "preview": text[:preview_chars],
            }
            if _has_page_break(el):
                info["contains_page_break"] = True
            if _has_section_break(el):
                info["contains_section_break"] = True
            if el.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None or \
               any(True for _ in el.iter(qn("w:numPr"))):
                info["numbered_or_bulleted"] = True
            blocks.append(info)
    return blocks


def _cover_page_hint(outline: list) -> dict:
    """Guess how many leading blocks form the cover page (up to and including
    the first page/section break)."""
    for b in outline:
        if b.get("contains_page_break") or b.get("contains_section_break"):
            return {"detected": True, "keep_blocks": b["index"] + 1,
                    "reason": "first page/section break at block index %d" % b["index"]}
    return {"detected": False, "keep_blocks": 0,
            "reason": "no explicit page/section break found in template body"}


_HEADING_HINTS = {1: ["heading 1", "titre 1", "title 1", "berschrift 1"],
                  2: ["heading 2", "titre 2", "berschrift 2"],
                  3: ["heading 3", "titre 3", "berschrift 3"],
                  4: ["heading 4", "titre 4", "berschrift 4"],
                  5: ["heading 5", "titre 5", "berschrift 5"],
                  6: ["heading 6", "titre 6", "berschrift 6"]}


def _suggest_style_map(doc) -> dict:
    idx = _style_index(doc)
    para_names = [s.name for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH]
    table_names = [s.name for s in doc.styles if s.type == WD_STYLE_TYPE.TABLE]

    def pick(candidates):
        for c in candidates:
            if c.lower() in idx:
                return idx[c.lower()].name
        return None

    sm = {}
    for lvl, hints in _HEADING_HINTS.items():
        found = pick(hints)
        if not found:  # fuzzy: any paragraph style ending with the level digit + 'heading'-ish
            for n in para_names:
                ln = n.lower()
                if ln.endswith(str(lvl)) and any(w in ln for w in ("head", "titre", "title", "chap")):
                    found = n
                    break
        if found:
            sm["h%d" % lvl] = found
    sm["paragraph"] = pick(["Normal", "Body Text"]) or "Normal"
    sm["bullet_list"] = pick(["List Bullet", "List Paragraph", "Liste à puces"])
    sm["ordered_list"] = pick(["List Number", "List Paragraph", "Liste numérotée"])
    sm["quote"] = pick(["Quote", "Intense Quote", "Citation"])
    sm["code_block"] = pick(["Code", "Source Code", "HTML Preformatted", "Code Block", "macro"])
    sm["caption"] = pick(["Caption", "Légende"])
    sm["hyperlink_char_style"] = pick(["Hyperlink", "Lien hypertexte"])
    # table style: prefer the style used by the first table in the template body
    tbl_style = None
    if doc.tables:
        tbl_style = _safe(lambda: doc.tables[0].style.name)
    if not tbl_style:
        for cand in ("Table Grid", "Grid Table 4 Accent 1", "Light Grid Accent 1"):
            if cand.lower() in idx:
                tbl_style = idx[cand.lower()].name
                break
    if not tbl_style and table_names:
        tbl_style = table_names[0]
    sm["table"] = tbl_style
    return {k: v for k, v in sm.items() if v}


def inspect_template(template_path: str) -> dict:
    if not os.path.isfile(template_path):
        raise FileNotFoundError("Template not found: %s" % template_path)
    doc = Document(template_path)

    styles = {"paragraph": [], "character": [], "table": []}
    for s in doc.styles:
        try:
            entry = {
                "name": s.name,
                "style_id": s.style_id,
                "builtin": s.builtin,
                "base_style": _safe(lambda: s.base_style.name if s.base_style else None),
            }
            if s.type == WD_STYLE_TYPE.PARAGRAPH:
                entry["font"] = _font_info(s.font)
                entry["alignment"] = _safe(
                    lambda: str(s.paragraph_format.alignment).split(".")[-1].split(" ")[0]
                    if s.paragraph_format.alignment is not None else None)
                styles["paragraph"].append(entry)
            elif s.type == WD_STYLE_TYPE.CHARACTER:
                entry["font"] = _font_info(s.font)
                styles["character"].append(entry)
            elif s.type == WD_STYLE_TYPE.TABLE:
                styles["table"].append(entry)
        except Exception:
            continue

    outline = _body_outline(doc)
    return {
        "template_path": os.path.abspath(template_path),
        "theme": _theme_info(template_path),
        "defaults": _doc_defaults(doc),
        "sections": _sections_info(doc),
        "styles": styles,
        "body_blocks": outline,
        "body_block_count": len(outline),
        "cover_page_hint": _cover_page_hint(outline),
        "placeholders": _collect_placeholders(doc),
        "suggested_style_map": _suggest_style_map(doc),
    }


# ---------------------------------------------------------------------------
# Placeholder replacement (formatting-preserving, span-aware)
# ---------------------------------------------------------------------------


def _replace_placeholders_in_paragraph(p, mapping: dict) -> int:
    """Replace {{key}} tokens in a paragraph even when a token is split across
    several runs. The replacement text takes the formatting of the first run
    the token starts in; untouched runs keep their formatting."""
    runs = p.runs
    if not runs:
        return 0
    full = "".join(r.text or "" for r in runs)
    matches = [m for m in PLACEHOLDER_RE.finditer(full) if m.group(1) in mapping]
    if not matches:
        return 0
    texts = [r.text or "" for r in runs]
    for m in reversed(matches):
        s, e = m.span()
        value = str(mapping[m.group(1)])
        pos = 0
        new_texts = []
        placed = False
        for t in texts:
            t_start, t_end = pos, pos + len(t)
            if t_end <= s or t_start >= e:
                new_texts.append(t)               # untouched run
            else:
                left = t[: max(0, s - t_start)]
                right = t[max(0, min(len(t), e - t_start)):]
                if not placed:
                    new_texts.append(left + value + right)
                    placed = True
                else:
                    new_texts.append(left + right)
            pos = t_end
        texts = new_texts
    for r, t in zip(runs, texts):
        if (r.text or "") != t:
            r.text = t
    return len(matches)


def _replace_placeholders_everywhere(doc, mapping: dict, scope_blocks=None) -> int:
    """scope_blocks: list of body block elements to restrict body replacement to
    (headers/footers are always processed)."""
    count = 0

    def do_paragraphs(paragraphs):
        nonlocal count
        for p in paragraphs:
            count += _replace_placeholders_in_paragraph(p, mapping)

    def do_tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    do_paragraphs(cell.paragraphs)
                    do_tables(cell.tables)

    from docx.text.paragraph import Paragraph
    from docx.table import Table
    if scope_blocks is None:
        do_paragraphs(doc.paragraphs)
        do_tables(doc.tables)
    else:
        for el in scope_blocks:
            if el.tag == qn("w:p"):
                do_paragraphs([Paragraph(el, doc)])
            elif el.tag == qn("w:tbl"):
                do_tables([Table(el, doc)])
    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part is not None:
                do_paragraphs(part.paragraphs)
                do_tables(part.tables)
    return count


# ---------------------------------------------------------------------------
# 2. MARKDOWN -> DOCX RENDERING
# ---------------------------------------------------------------------------

DEFAULT_STYLE_CANDIDATES = {
    "h1": ["Heading 1", "Titre 1"],
    "h2": ["Heading 2", "Titre 2"],
    "h3": ["Heading 3", "Titre 3"],
    "h4": ["Heading 4", "Titre 4"],
    "h5": ["Heading 5", "Titre 5"],
    "h6": ["Heading 6", "Titre 6"],
    "paragraph": ["Normal"],
    "bullet_list": ["List Bullet", "List Paragraph"],
    "ordered_list": ["List Number", "List Paragraph"],
    "quote": ["Quote", "Intense Quote", "Citation"],
    "code_block": ["Code", "Source Code", "HTML Preformatted", "macro"],
    "caption": ["Caption", "Légende"],
    "table": [],  # resolved from template's own tables / Table Grid
    "hyperlink_char_style": ["Hyperlink", "Lien hypertexte"],
}


class _StyleResolver:
    def __init__(self, doc, style_map: Optional[dict], warnings: list,
                 template_table_style: Optional[str] = None):
        self.doc = doc
        self.idx = _style_index(doc)
        self.override = {k.lower(): v for k, v in (style_map or {}).items()}
        self.warnings = warnings
        self.cache: dict = {}
        self.template_table_style = template_table_style

    def resolve(self, key: str) -> Optional[str]:
        """Return an existing template style NAME for a semantic key, or None."""
        k = key.lower()
        if k in self.cache:
            return self.cache[k]
        result = None
        if k in self.override:
            wanted = self.override[k]
            st = self.idx.get(str(wanted).lower())
            if st is not None:
                result = st.name
            else:
                self.warnings.append(
                    "style_map['%s'] = '%s' does not exist in the template; "
                    "falling back to defaults." % (key, wanted))
        if result is None:
            for cand in DEFAULT_STYLE_CANDIDATES.get(k, []):
                st = self.idx.get(cand.lower())
                if st is not None:
                    result = st.name
                    break
        if result is None and k.startswith("h") and k[1:].isdigit():
            self.warnings.append(
                "No style found for '%s'; using 'Normal' with bold." % key)
        self.cache[k] = result
        return result

    def resolve_table(self) -> Optional[str]:
        if "table" in self.cache:
            return self.cache["table"]
        result = None
        if "table" in self.override:
            st = self.idx.get(str(self.override["table"]).lower())
            if st is not None and st.type == WD_STYLE_TYPE.TABLE:
                result = st.name
            else:
                self.warnings.append(
                    "style_map['table'] = '%s' is not a table style of the "
                    "template; falling back." % self.override["table"])
        if result is None and self.template_table_style:
            st = self.idx.get(self.template_table_style.lower())
            if st is not None:
                result = st.name
        if result is None:
            for cand in ("Table Grid", "Grid Table 4 Accent 1"):
                st = self.idx.get(cand.lower())
                if st is not None:
                    result = st.name
                    break
        self.cache["table"] = result
        return result


def _add_hyperlink(paragraph, url: str, resolver: _StyleResolver):
    """Create a real w:hyperlink and return the run element container to fill."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    paragraph._p.append(hl)
    return hl


def _new_run_in(container, paragraph):
    """Append a w:r inside `container` (a w:p or w:hyperlink) and wrap it."""
    from docx.text.run import Run
    r = OxmlElement("w:r")
    container.append(r)
    return Run(r, paragraph)


class _MarkdownRenderer:
    def __init__(self, doc, resolver: _StyleResolver, warnings: list,
                 base_dir: Optional[str]):
        self.doc = doc
        self.res = resolver
        self.warnings = warnings
        self.base_dir = base_dir or os.getcwd()
        self.used_styles: set = set()
        self.counts = {"headings": 0, "paragraphs": 0, "tables": 0,
                       "lists": 0, "images": 0, "code_blocks": 0}

    # -------------------------------------------------- paragraph factory
    def _p(self, style_key: Optional[str], fallback_bold: bool = False):
        name = self.res.resolve(style_key) if style_key else None
        p = self.doc.add_paragraph()
        if name:
            try:
                p.style = self.doc.styles[name]
                self.used_styles.add(name)
            except Exception:
                name = None
        if name is None and fallback_bold:
            p._fallback_bold = True  # inline renderer will bold the runs
        return p

    # -------------------------------------------------- inline rendering
    def _inline(self, paragraph, children, container=None, fmt=None):
        fmt = fmt or {}
        container = container if container is not None else paragraph._p
        for tok in children or []:
            t = tok.get("type")
            if t == "text":
                self._emit_text(paragraph, container, tok.get("raw", ""), fmt)
            elif t == "strong":
                self._inline(paragraph, tok.get("children"), container,
                             {**fmt, "bold": True})
            elif t == "emphasis":
                self._inline(paragraph, tok.get("children"), container,
                             {**fmt, "italic": True})
            elif t == "strikethrough":
                self._inline(paragraph, tok.get("children"), container,
                             {**fmt, "strike": True})
            elif t == "codespan":
                self._emit_text(paragraph, container, tok.get("raw", ""),
                                {**fmt, "code": True})
            elif t in ("linebreak", "softbreak"):
                run = _new_run_in(container, paragraph)
                if t == "linebreak":
                    run.add_break()
                else:
                    run.text = " "
            elif t == "link":
                url = tok.get("attrs", {}).get("url", "")
                hl = _add_hyperlink(paragraph, url, self.res)
                char_style = self.res.resolve("hyperlink_char_style")
                self._inline(paragraph, tok.get("children"), hl,
                             {**fmt, "_hyperlink_style": char_style})
            elif t == "image":
                self._emit_image(paragraph, tok)
            elif t == "inline_html":
                self._emit_text(paragraph, container, tok.get("raw", ""), fmt)
            elif tok.get("children"):
                self._inline(paragraph, tok["children"], container, fmt)
            elif tok.get("raw"):
                self._emit_text(paragraph, container, tok["raw"], fmt)

    def _emit_text(self, paragraph, container, text, fmt):
        if text == "":
            return
        run = _new_run_in(container, paragraph)
        run.text = text
        if fmt.get("bold") or getattr(paragraph, "_fallback_bold", False):
            run.bold = True
        if fmt.get("italic"):
            run.italic = True
        if fmt.get("strike"):
            run.font.strike = True
        if fmt.get("code"):
            run.font.name = "Consolas"
            rPr = run._r.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn("w:cs"), "Consolas")
        hs = fmt.get("_hyperlink_style")
        if hs:
            try:
                run.style = self.doc.styles[hs]
                self.used_styles.add(hs)
            except Exception:
                run.font.underline = True
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        elif fmt.get("_hyperlink_style") is None and fmt.get("_in_link"):
            run.font.underline = True

    def _emit_image(self, paragraph, tok):
        url = tok.get("attrs", {}).get("url", "")
        alt = "".join(c.get("raw", "") for c in tok.get("children") or [])
        path = url
        if not os.path.isabs(path):
            path = os.path.join(self.base_dir, url)
        if url.startswith(("http://", "https://")) or not os.path.isfile(path):
            self.warnings.append("Image not embedded (missing or remote): %s" % url)
            run = _new_run_in(paragraph._p, paragraph)
            run.text = "[image: %s]" % (alt or url)
            run.italic = True
            return
        run = paragraph.add_run()
        section = self.doc.sections[-1]
        text_width = section.page_width - section.left_margin - section.right_margin
        try:
            pic = run.add_picture(path)
            if pic.width > text_width:
                ratio = text_width / pic.width
                pic.width = Emu(int(pic.width * ratio))
                pic.height = Emu(int(pic.height * ratio))
            self.counts["images"] += 1
        except Exception as e:
            self.warnings.append("Failed to embed image %s: %s" % (url, e))

    # -------------------------------------------------- block rendering
    def render(self, tokens):
        for tok in tokens:
            t = tok.get("type")
            if t == "heading":
                lvl = tok.get("attrs", {}).get("level", 1)
                p = self._p("h%d" % min(lvl, 6), fallback_bold=True)
                self._inline(p, tok.get("children"))
                self.counts["headings"] += 1
            elif t == "paragraph":
                p = self._p("paragraph")
                self._inline(p, tok.get("children"))
                self.counts["paragraphs"] += 1
            elif t == "list":
                self._render_list(tok, level=0)
                self.counts["lists"] += 1
            elif t == "block_quote":
                self._render_blockquote(tok)
            elif t == "block_code":
                self._render_code(tok)
            elif t == "thematic_break":
                self._render_hr()
            elif t == "table":
                self._render_table(tok)
            elif t in ("blank_line",):
                continue
            elif t == "block_html":
                p = self._p("paragraph")
                p.add_run(tok.get("raw", ""))
            elif tok.get("children"):
                self.render(tok["children"])

    def _list_style(self, ordered: bool, level: int) -> Optional[str]:
        base_key = "ordered_list" if ordered else "bullet_list"
        base = self.res.resolve(base_key)
        if base and level > 0:
            # try 'List Bullet 2' / 'List Number 3' etc.
            candidate = "%s %d" % (base, level + 1)
            st = self.res.idx.get(candidate.lower())
            if st is not None:
                return st.name
        return base

    def _render_list(self, tok, level: int):
        ordered = tok.get("attrs", {}).get("ordered", False)
        for item in tok.get("children", []):
            first_para_done = False
            for child in item.get("children", []):
                ct = child.get("type")
                if ct in ("block_text", "paragraph"):
                    style = self._list_style(ordered, level)
                    p = self.doc.add_paragraph()
                    if style:
                        try:
                            p.style = self.doc.styles[style]
                            self.used_styles.add(style)
                        except Exception:
                            style = None
                    if style is None:
                        p.paragraph_format.left_indent = Pt(18 * (level + 1))
                        pref = "%d. " % 1 if ordered else "• "
                        p.add_run(pref)
                    self._inline(p, child.get("children"))
                    first_para_done = True
                elif ct == "list":
                    self._render_list(child, level + 1)
                else:
                    self.render([child])

    def _render_blockquote(self, tok):
        style = self.res.resolve("quote")
        for child in tok.get("children", []):
            if child.get("type") == "paragraph":
                p = self.doc.add_paragraph()
                if style:
                    try:
                        p.style = self.doc.styles[style]
                        self.used_styles.add(style)
                    except Exception:
                        style = None
                if style is None:
                    p.paragraph_format.left_indent = Pt(24)
                self._inline(p, child.get("children"))
                if style is None:
                    for r in p.runs:
                        r.italic = True
            else:
                self.render([child])

    def _render_code(self, tok):
        style = self.res.resolve("code_block")
        raw = tok.get("raw", "").rstrip("\n")
        p = self.doc.add_paragraph()
        if style:
            try:
                p.style = self.doc.styles[style]
                self.used_styles.add(style)
            except Exception:
                style = None
        run = p.add_run(raw)
        if style is None:
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F2F2F2")
            pPr.append(shd)
        self.counts["code_blocks"] += 1

    def _render_hr(self):
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "auto")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _render_table(self, tok):
        head = next((c for c in tok.get("children", [])
                     if c.get("type") == "table_head"), None)
        body = next((c for c in tok.get("children", [])
                     if c.get("type") == "table_body"), None)
        head_cells = head.get("children", []) if head else []
        body_rows = body.get("children", []) if body else []
        n_cols = len(head_cells) or (len(body_rows[0].get("children", []))
                                     if body_rows else 0)
        if n_cols == 0:
            return
        n_rows = (1 if head_cells else 0) + len(body_rows)
        table = self.doc.add_table(rows=n_rows, cols=n_cols)
        style = self.res.resolve_table()
        if style:
            try:
                table.style = self.doc.styles[style]
                self.used_styles.add(style)
            except Exception:
                pass
        r = 0
        if head_cells:
            for c_i, cell_tok in enumerate(head_cells):
                if c_i >= n_cols:
                    break
                cell = table.cell(0, c_i)
                p = cell.paragraphs[0]
                self._inline(p, cell_tok.get("children"))
                for run in p.runs:
                    if run.bold is None:
                        run.bold = True
            r = 1
        for row_tok in body_rows:
            for c_i, cell_tok in enumerate(row_tok.get("children", [])):
                if c_i >= n_cols:
                    break
                p = table.cell(r, c_i).paragraphs[0]
                self._inline(p, cell_tok.get("children"))
            r += 1
        self.counts["tables"] += 1
        # spacer paragraph so consecutive tables don't merge
        self.doc.add_paragraph()


def _parse_markdown(md_text: str):
    md = mistune.create_markdown(
        renderer=None, plugins=["table", "strikethrough", "url"])
    return md(md_text)


def create_document(template_path: str,
                    output_path: str,
                    markdown_text: str,
                    keep_blocks: int = 0,
                    page_break_before_content: Optional[bool] = None,
                    placeholders: Optional[dict] = None,
                    style_map: Optional[dict] = None,
                    markdown_base_dir: Optional[str] = None) -> dict:
    """Render markdown into a copy of the template. See server.py for the
    full contract of each parameter."""
    if not os.path.isfile(template_path):
        raise FileNotFoundError("Template not found: %s" % template_path)
    warnings: list = []
    doc = Document(template_path)
    # Capture the table style used by the template's own tables BEFORE any
    # body blocks are removed, so generated tables inherit the same look.
    template_table_style = _safe(
        lambda: doc.tables[0].style.name if doc.tables else None)
    blocks = _iter_body_blocks(doc)
    total = len(blocks)

    if keep_blocks < -1:
        raise ValueError("keep_blocks must be >= -1")
    if keep_blocks > total:
        warnings.append("keep_blocks=%d exceeds template body block count (%d); "
                        "keeping all blocks." % (keep_blocks, total))
        keep_blocks = total

    if keep_blocks == -1:
        kept = blocks
    else:
        kept = blocks[:keep_blocks]
        for el in blocks[keep_blocks:]:
            el.getparent().remove(el)

    # Placeholder substitution on kept blocks + headers/footers
    replaced = 0
    if placeholders:
        replaced = _replace_placeholders_everywhere(doc, placeholders,
                                                    scope_blocks=None)

    # Page break between kept cover content and new report content
    if page_break_before_content is None:
        page_break_before_content = bool(kept) and not (
            kept and (_has_page_break(kept[-1]) or _has_section_break(kept[-1])))
    if kept and page_break_before_content:
        from docx.enum.text import WD_BREAK
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    resolver = _StyleResolver(doc, style_map, warnings,
                              template_table_style=template_table_style)
    renderer = _MarkdownRenderer(doc, resolver, warnings,
                                 base_dir=markdown_base_dir)
    tokens = _parse_markdown(markdown_text)
    renderer.render(tokens)

    # Ensure the body-level sectPr stays the LAST child of w:body
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    if sectPr is not None:
        body.append(sectPr)

    out_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)

    return {
        "output_path": os.path.abspath(output_path),
        "template_path": os.path.abspath(template_path),
        "kept_template_blocks": len(kept) if keep_blocks != -1 else total,
        "placeholders_replaced": replaced,
        "elements_rendered": renderer.counts,
        "template_styles_used": sorted(renderer.used_styles),
        "warnings": warnings,
    }


# ---------------------------------------------------------------------------
# 3. VALIDATION
# ---------------------------------------------------------------------------


def validate_document(document_path: str, template_path: str) -> dict:
    if not os.path.isfile(document_path):
        raise FileNotFoundError("Document not found: %s" % document_path)
    if not os.path.isfile(template_path):
        raise FileNotFoundError("Template not found: %s" % template_path)
    doc = Document(document_path)
    tpl = Document(template_path)

    tpl_styles = {s.name for s in tpl.styles} | {s.style_id for s in tpl.styles}
    tpl_fonts = set()
    for s in tpl.styles:
        n = _safe(lambda: s.font.name)
        if n:
            tpl_fonts.add(n)
    theme = _theme_info(template_path)
    tpl_fonts |= set(theme.get("fonts", {}).values())
    tpl_fonts |= {"Consolas", "Courier New"}  # accepted for code rendering

    used_para_styles: set = set()
    used_char_styles: set = set()
    foreign_fonts: dict = {}
    leftover_placeholders: set = set()
    n_paragraphs = n_tables = 0

    def scan_paragraph(p):
        nonlocal n_paragraphs
        n_paragraphs += 1
        used_para_styles.add(p.style.name)
        for m in PLACEHOLDER_RE.finditer(p.text):
            leftover_placeholders.add(m.group(1))
        for r in p.runs:
            if r.style and r.style.name != "Default Paragraph Font":
                used_char_styles.add(r.style.name)
            fname = _safe(lambda: r.font.name)
            if fname and fname not in tpl_fonts:
                foreign_fonts[fname] = foreign_fonts.get(fname, 0) + 1

    def scan_tables(tables):
        nonlocal n_tables
        for t in tables:
            n_tables += 1
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        scan_paragraph(p)
                    scan_tables(cell.tables)

    for p in doc.paragraphs:
        scan_paragraph(p)
    scan_tables(doc.tables)

    unknown_para = sorted(s for s in used_para_styles if s not in tpl_styles)
    unknown_char = sorted(s for s in used_char_styles if s not in tpl_styles)

    issues = []
    if unknown_para:
        issues.append("Paragraph styles not defined in template: %s" % unknown_para)
    if unknown_char:
        issues.append("Character styles not defined in template: %s" % unknown_char)
    if leftover_placeholders:
        issues.append("Unreplaced placeholders remain: %s"
                      % sorted(leftover_placeholders))
    if foreign_fonts:
        issues.append("Fonts not part of the template appear in runs: %s"
                      % foreign_fonts)
    return {
        "document_path": os.path.abspath(document_path),
        "template_path": os.path.abspath(template_path),
        "conforms": not (unknown_para or unknown_char or leftover_placeholders),
        "issues": issues,
        "stats": {
            "paragraphs": n_paragraphs,
            "tables": n_tables,
            "paragraph_styles_used": sorted(used_para_styles),
            "character_styles_used": sorted(used_char_styles),
        },
    }
